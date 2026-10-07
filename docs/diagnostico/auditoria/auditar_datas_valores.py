import openpyxl, sys, re, datetime as dt, json
from dateutil.relativedelta import relativedelta
from openpyxl.utils import get_column_letter as L
p=sys.argv[1]
ws=openpyxl.load_workbook(p).worksheets[0]
wv=openpyxl.load_workbook(p,data_only=True).worksheets[0]
REF=dt.datetime(2026,10,7)
cols={'A':'secretaria','B':'numero','C':'contratada','D':'objeto','E':'unidade','G':'quant','H':'H','I':'v_unit','J':'v_mensal','K':'v_anual','L':'v_total','M':'proc_mae','N':'fat2023','O':'fat2024','P':'fat2025','Q':'fat2026','R':'inicio','S':'prazo','T':'termino','U':'doc_base','V':'inicio_vig_ref','W':'garantia','X':'obs','Y':'apost_situacao','Z':'portaria','AA':'fiscal_pres','AB':'fiscal1','AC':'fiscal2','AD':'fiscal_sub','AE':'gestor','AF':'gestor_sub','AG':'dt_assin','AI':'dt_doerj','AK':'dt_pncp','AM':'siafe'}
rows=[]
for r in list(range(12,21))+list(range(22,53)):
    d={'row':r}
    for c,n in cols.items():
        col=openpyxl.utils.column_index_from_string(c)
        d[n]=wv.cell(r,col).value; d[n+'_f']=ws.cell(r,col).value
    rows.append(d)
def isnum(v): return isinstance(v,(int,float)) and not isinstance(v,bool)
def placeholder(v): return isinstance(v,str) and (re.fullmatch(r'[\s\*xX]+',v) or v.strip() in ('N/C','* * *','**','XXX'))
def empty(v): return v is None or (isinstance(v,str) and v.strip()=='')
out=[]
for d in rows:
    if empty(d['numero']) and empty(d['inicio']): d['tipo']='sem_numero'
    issues=[]
    # dates
    if isinstance(d['inicio'],dt.datetime) and isnum(d['prazo']):
        exp=d['inicio']+relativedelta(months=int(d['prazo']))-dt.timedelta(days=1)
        d['termino_esperado']=exp
        if isinstance(d['termino'],dt.datetime):
            diff=(d['termino']-exp).days
            d['dif_dias']=diff
    t=d['termino'] if isinstance(d['termino'],dt.datetime) else None
    d['dias_para_vencer']=(t-REF).days if t else None
    # values
    if isnum(d['v_mensal']) and isnum(d['v_total']) and isnum(d['prazo']):
        d['total_esperado_mensal_x_prazo']=round(d['v_mensal']*d['prazo'],2)
    if isnum(d['quant']) and isnum(d['v_unit']):
        d['q_x_u']=round(d['quant']*d['v_unit'],2)
    out.append(d)
for d in out:
    print(d['row'], d['secretaria'], repr(d['numero']), '|', (d['contratada'] or '').strip(),
          '| ini',d['inicio'] and d['inicio'].date(), 'prazo',d['prazo'],'term',d['termino'] and getattr(d['termino'],'date',lambda:d['termino'])(),
          '| T_f',d['termino_f'] if isinstance(d['termino_f'],str) else 'hard',
          '| esp',d.get('termino_esperado') and d['termino_esperado'].date(),'dif',d.get('dif_dias'),'| d2v',d['dias_para_vencer'])
print()
for d in out:
    print(d['row'],(d['contratada'] or '').strip()[:18],'| Q',d['quant'],'U',d['v_unit'],'M',d['v_mensal'],'A',d['v_anual'],'T',d['v_total'],'| qxu',d.get('q_x_u'),'| mxprazo',d.get('total_esperado_mensal_x_prazo'),
          '| ratio T/(M*prazo)', round(d['v_total']/d['total_esperado_mensal_x_prazo'],3) if d.get('total_esperado_mensal_x_prazo') else None)
json.dump(out, open(sys.argv[2],'w'), default=str, ensure_ascii=False, indent=1)
