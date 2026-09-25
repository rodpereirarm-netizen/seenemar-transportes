"""Matriz de atividades por ator — GT PROPAG (posição 25/09/2026).
Gera matriz_atividades_atores.html (convertido em PDF) e Matriz_Atividades_Atores_GT_PROPAG.xlsx.
Fontes: F1 = documento “GT FAETEC — Macroprocesso e Equipe” (equipe por etapa e função de cada integrante);
F2 = sistema em produção (consulta somente leitura em 25/09/2026). Nada é gravado no sistema."""
import html, openpyxl
from pathlib import Path
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
D = Path(__file__).parent
e = html.escape

# ------------------------------------------------------------------ atores
# código: (nome do ator, quem é, grupo)
ATORES = {
 'COORD':  ('Coordenação do GT', 'Karina Ferrarez (SEDES) · Vinicius Murat (FAETEC)', 'GT · Frente I'),
 'APROV':  ('Aprovação dos artefatos', 'Cristiane Vaz dos Santos Aguiar (SECTI)', 'GT · Frente I'),
 'CONF':   ('Conformidade dos artefatos', 'Elias Conceição Magalhães (SEDES)', 'GT · Frente I'),
 'ELAB':   ('Elaboração de artefatos', 'Andressa, Allana, Taina (SEDES) · Thamyres (SECTI) · Mara, Thailane, Marco (PRODERJ)', 'GT · Frente II'),
 'ATAS':   ('Pesquisa de atas de TI', 'Pascoal (SECTI)', 'GT · Frente II'),
 'PCA':    ('Ponto focal PCA', 'Jhonatan Silva Santos (FAETEC)', 'GT · Frente III'),
 'TI':     ('Divisão de TI', 'Gibson Cruz da Silva (FAETEC)', 'GT · Frente III'),
 'DOCS':   ('Ponto focal documentos', "Luene Fernandes Curvello d'Ávila (FAETEC)", 'GT · Frente III'),
 'DEMAND': ('Área demandante', 'Diretoria/área técnica da FAETEC que pede a contratação', 'FAETEC'),
 'DGI':    ('FAETEC/DGI', 'Diretoria citada na equipe das Etapas I e II', 'FAETEC'),
 'DIVPL':  ('FAETEC/Div. Planejamento', 'Divisão citada na equipe da Etapa II', 'FAETEC'),
 'VPA':    ('FAETEC/VPA', 'Instância que aprova e autoriza (Etapa I)', 'FAETEC'),
 'FAETEC': ('FAETEC · ponto focal a nomear', 'Etapas IV e V: “FAETEC (100%)”, sem ponto focal na Relação de Participantes', 'FAETEC'),
 'ORDEN':  ('Ordenador de despesas', 'Autoridade da FAETEC', 'FAETEC'),
 'JUR':    ('Assessoria jurídica', 'Análise jurídica da minuta', 'FAETEC'),
 'UCI':    ('Controle interno (UCI)', 'Nota de auditoria', 'FAETEC'),
 'GEREN':  ('Órgão gerenciador da ata', 'Ex.: SEAD/MA, CELIC/RS, TJMA, SEPLAG/MG, PRODERJ', 'Externo'),
 'FORN':   ('Fornecedor', 'Detentor da ata / contratado', 'Externo'),
 'SEPLAG': ('SEPLAG', 'Comunicação prévia · 15 dias corridos', 'Externo'),
 'PRODERJ':('PRODERJ (análise TIC)', 'Análise técnica · 20 dias úteis (+20)', 'Externo'),
 'CGE':    ('CGE', 'Análise · 15 dias corridos', 'Externo'),
}

# Papel na atividade: E = executa (quem faz) · A = apoia/fornece insumo · V = valida/aprova · X = órgão externo que responde
# Base: DOC = atribuição expressa no documento F1 · FUN = derivada da função do integrante em F1 · NAT = natureza do ato/norma · PROP = proposta (sem definição no F1)
ATIV = [
 # Etapa I · Planejamento da contratação — equipe F1: FAETEC/DGI · GT · FAETEC/VPA · executor: Elaboração (Frente II)
 (1, 1, 'Identificação da demanda', {'DEMAND': 'E', 'DGI': 'A', 'COORD': 'A'}, 'Justificativa da necessidade', '', 'NAT', 'A área que precisa formaliza a justificativa; a coordenação recebe e dá entrada no GT.'),
 (1, 2, 'Inclusão no PCA / PEDTIC', {'PCA': 'E', 'TI': 'A', 'CONF': 'V'}, 'Item do PCA/PEDTIC confirmado', 'Ponto de atenção: fora do PEDTIC exige justificativa técnica e revisão extraordinária', 'FUN', 'Jhonatan é o “Ponto Focal PCA” no F1. No sistema, a atividade está com Luene em 1 demanda.'),
 (1, 3, 'Elaboração do DOD', {'ELAB': 'E', 'DEMAND': 'A', 'CONF': 'V', 'APROV': 'V'}, 'DOD', '', 'DOC', 'Executor da Etapa I no F1: elaboração de artefatos (Frente II).'),
 (1, 4, 'Elaboração do ETP', {'ELAB': 'E', 'TI': 'A', 'CONF': 'V', 'APROV': 'V'}, 'ETP + manifestação do TI', 'Participante de RP reaproveita o ETP do gerenciador', 'DOC', 'O TI (Gibson) dá a manifestação técnica; no sistema, a atividade está com ele em 2 demandas.'),
 (1, 5, 'Elaboração do TR', {'ELAB': 'E', 'DEMAND': 'A', 'TI': 'A', 'CONF': 'V', 'APROV': 'V'}, 'Termo de Referência', '', 'DOC', ''),
 (1, 6, 'Mapa de Riscos', {'ELAB': 'E', 'CONF': 'V', 'APROV': 'V'}, 'Mapa de riscos', '', 'DOC', ''),
 (1, 7, 'Aprovação e autorização (VPA)', {'VPA': 'E', 'COORD': 'A', 'DOCS': 'A'}, 'Autorização para prosseguir', '', 'DOC', 'VPA está na equipe da Etapa I no F1; coordenação e ponto focal de documentos instruem e encaminham (proposta).'),
 # Etapa II · Adesão à ARP — equipe F1: GT · FAETEC/Div. Planejamento · FAETEC/DGI · executor: Elaboração (Frente II)
 (2, 1, 'Anuência do órgão gerenciador', {'ELAB': 'E', 'DIVPL': 'A', 'GEREN': 'X'}, 'Ofício de anuência do gerenciador', '', 'DOC', 'O GT prepara o pedido; o gerenciador responde.'),
 (2, 2, 'Anuência do fornecedor', {'ELAB': 'E', 'DIVPL': 'A', 'FORN': 'X'}, 'Aceite do fornecedor', '', 'DOC', ''),
 (2, 3, 'Verificação de ata válida', {'ATAS': 'E', 'ELAB': 'A', 'CONF': 'V'}, 'Ata vigente, com saldo e adesão possível', 'Gate: sem ata válida, o fluxo de ARP é interrompido', 'FUN', 'Pascoal tem a função “Pesquisa atas TI” no F1.'),
 (2, 4, 'Inserção dos documentos da ARP', {'DOCS': 'E', 'DIVPL': 'A'}, 'Documentos da ata no processo SEI', '', 'FUN', 'Luene é “Ponto Focal Documentos” no F1.'),
 (2, 5, 'Compatibilidade do objeto', {'ELAB': 'E', 'TI': 'A', 'DEMAND': 'A', 'APROV': 'V'}, 'Análise de compatibilidade', '', 'DOC', ''),
 (2, 6, 'Justificativa de vantajosidade', {'ELAB': 'E', 'CONF': 'V', 'APROV': 'V'}, 'Justificativa (Enunciado 27 PGE-RJ)', '', 'DOC', ''),
 # Etapa III · Comunicações e análises — equipe F1: FAETEC · executor: Ponto focal FAETEC (Frente III)
 (3, 1, 'Comunicação à SEPLAG', {'DOCS': 'E', 'SEPLAG': 'X'}, 'Comunicação protocolada', 'Prazo legal: 15 dias corridos (Dec. 48.821/2023, art. 2º)', 'PROP', 'O F1 atribui a etapa ao ponto focal FAETEC (Frente III), sem dizer qual dos três; proposta: Luene (documentos).'),
 (3, 2, 'Análise técnica do PRODERJ (TIC)', {'TI': 'E', 'PRODERJ': 'X'}, 'Parecer técnico do PRODERJ', 'Prazo legal: 20 dias úteis, prorrogáveis por +20 (IN PRODERJ/PRE 05/2024, art. 4º)', 'PROP', 'Proposta: Gibson (Divisão de TI) envia e acompanha.'),
 (3, 3, 'Registro no Contratos.gov.br', {'PCA': 'E'}, 'Registro no sistema', '', 'PROP', 'Os três pontos focais têm “Lançamento dados Sistemas” no F1; proposta: Jhonatan.'),
 # Etapa IV · Pesquisa de preços — equipe F1: FAETEC (100%) · sem ponto focal nomeado
 (4, 1, 'Pesquisa de preços no SIGA', {'FAETEC': 'E'}, 'Pesquisa registrada no SIGA', 'Sem ponto focal nomeado', 'DOC', 'O F1 diz apenas “FAETEC (100%)”.'),
 (4, 2, 'Relatório Analítico (RAPP)', {'FAETEC': 'E'}, 'RAPP', 'Sem ponto focal nomeado', 'DOC', ''),
 (4, 3, 'Ateste do setor técnico', {'TI': 'E', 'FAETEC': 'A'}, 'Ateste técnico', '', 'PROP', 'As 20 demandas são de TI; proposta: Divisão de TI atesta.'),
 (4, 4, 'Checklist PGE', {'FAETEC': 'E'}, 'Checklist (Enunciado 27 PGE-RJ)', 'Sem ponto focal nomeado', 'DOC', ''),
 (4, 5, 'Aprovação da pesquisa (SIGA)', {'FAETEC': 'V'}, 'Pesquisa aprovada no SIGA', 'Sem ponto focal nomeado', 'DOC', 'Aprovação pela autoridade competente da FAETEC.'),
 (4, 6, 'Disponibilidade orçamentária', {'FAETEC': 'E'}, 'Declaração de disponibilidade', 'Sem ponto focal nomeado', 'NAT', 'Setor orçamentário da FAETEC.'),
 (4, 7, 'Declaração do ordenador de despesas', {'ORDEN': 'E'}, 'Declaração assinada', '', 'NAT', ''),
 # Etapa V · Formalização e publicidade — equipe F1: FAETEC (100%) · sem ponto focal nomeado
 (5, 1, 'Documentos de habilitação', {'FAETEC': 'E', 'FORN': 'X'}, 'Habilitação conferida', 'Sem ponto focal nomeado', 'NAT', 'O fornecedor entrega; a FAETEC confere.'),
 (5, 2, 'Minuta do contrato', {'FAETEC': 'E'}, 'Minuta', 'Sem ponto focal nomeado', 'DOC', ''),
 (5, 3, 'Análise jurídica', {'JUR': 'E'}, 'Parecer jurídico', '', 'NAT', ''),
 (5, 4, 'Nota de auditoria (UCI)', {'UCI': 'E'}, 'Nota de auditoria', '', 'NAT', ''),
 (5, 5, 'Análise da CGE', {'FAETEC': 'E', 'CGE': 'X'}, 'Manifestação da CGE', 'Prazo legal: 15 dias corridos (Dec. 48.821/2023)', 'NAT', 'A FAETEC envia; a CGE responde.'),
 (5, 6, 'Empenho', {'FAETEC': 'E', 'ORDEN': 'V'}, 'Nota de empenho', 'Sem ponto focal nomeado', 'NAT', 'Setor financeiro da FAETEC.'),
 (5, 7, 'Assinatura do contrato', {'ORDEN': 'E', 'FORN': 'X'}, 'Contrato assinado', '', 'NAT', ''),
 (5, 8, 'Publicação (D.O. / PNCP)', {'FAETEC': 'E'}, 'Extrato publicado no D.O. e no PNCP', 'Sem ponto focal nomeado', 'NAT', 'Lei 14.133/2021.'),
 (5, 9, 'Cadastro e-TCE (SIGFIS)', {'FAETEC': 'E'}, 'Cadastro no SIGFIS', 'Sem ponto focal nomeado', 'NAT', ''),
]
assert len(ATIV) == 32 and [sum(1 for a in ATIV if a[0] == k) for k in range(1, 6)] == [7, 6, 3, 7, 9]

ETAPA = {1: ('I', 'Planejamento da contratação', 'FAETEC/DGI · GT · FAETEC/VPA', 'Elaboração de artefatos · Frente II'),
         2: ('II', 'Adesão à ARP', 'GT · FAETEC/Div. Planejamento · FAETEC/DGI', 'Elaboração de artefatos · Frente II'),
         3: ('III', 'Comunicações e análises', 'FAETEC', 'Ponto focal FAETEC · Frente III'),
         4: ('IV', 'Pesquisa de preços', 'FAETEC (100%)', 'Ponto focal FAETEC — não nomeado'),
         5: ('V', 'Formalização e publicidade', 'FAETEC (100%)', 'Ponto focal FAETEC — não nomeado')}
# Instâncias de cada atividade por etapa: sistema hoje (17 demandas) e projeção 2ª carga + Tela (21 demandas)
INST_HOJE = {1: 17, 2: 4, 3: 7, 4: 7, 5: 7}
INST_PROJ = {1: 21, 2: 6, 3: 9, 4: 9, 5: 9}
assert sum(INST_HOJE[a[0]] for a in ATIV) == 276 and sum(INST_PROJ[a[0]] for a in ATIV) == 354
# Situação no sistema (consulta somente leitura, 25/09/2026): atividades com responsável designado
NO_SISTEMA = {(1, 2): "1 de 17 · Luene", (1, 4): "2 de 17 · Gibson"}

PAPEL = {'E': 'Executa', 'A': 'Apoia', 'V': 'Valida/aprova', 'X': 'Responde (externo)'}
BASE = {'DOC': 'Documento do GT', 'FUN': 'Função no documento', 'NAT': 'Natureza do ato', 'PROP': 'Proposta'}

def quem(p, papel):
    return [c for c, v in p.items() if v == papel]

# ------------------------------------------------------------------ carga por ator
carga = {}
for et, o, nome, p, *_ in ATIV:
    for c, v in p.items():
        d = carga.setdefault(c, {'E': [], 'A': [], 'V': [], 'X': [], 'hoje': 0, 'proj': 0})
        d[v].append(f'{ETAPA[et][0]}.{o} {nome}')
        if v in ('E', 'X'):
            d['hoje'] += INST_HOJE[et]; d['proj'] += INST_PROJ[et]
tot_e_hoje = sum(INST_HOJE[a[0]] for a in ATIV)

# ------------------------------------------------------------------ HTML
def cel(c, papel):
    n = ATORES[c][0]
    cls = {'E': 'e', 'A': 'a', 'V': 'v', 'X': 'x'}[papel]
    return f'<span class="tag {cls}">{e(n)}</span>'

linhas_ativ = ''
for et in range(1, 6):
    r, nome_et, eq, execu = ETAPA[et]
    linhas_ativ += f'<tr class="grp"><td colspan="6">Etapa {r} · {e(nome_et)} <span>· equipe no documento: {e(eq)} · executor: {e(execu)}</span></td></tr>'
    for (a_et, o, nome, p, entrega, alerta, base, obs) in [a for a in ATIV if a[0] == et]:
        ex = ' '.join(cel(c, 'E') for c in quem(p, 'E'))
        ap = ' '.join(cel(c, 'A') for c in quem(p, 'A')) or '—'
        va = ' '.join(cel(c, 'V') for c in quem(p, 'V')) or '—'
        xt = ' '.join(cel(c, 'X') for c in quem(p, 'X')) or '—'
        det = e(entrega) + (f'<br><b class="alerta">{e(alerta)}</b>' if alerta else '') + (f'<br><span class="obs">{e(obs)}</span>' if obs else '')
        sis = NO_SISTEMA.get((et, o), f'0 de {INST_HOJE[et]}')
        linhas_ativ += (f'<tr><td class="id">{r}.{o}</td><td><b>{e(nome)}</b><br><span class="obs">{det}</span></td>'
                        f'<td>{ex}</td><td>{ap}</td><td>{va}<br>{xt if xt != "—" else ""}</td>'
                        f'<td><span class="base {base.lower()}">{BASE[base]}</span><br><span class="obs">no sistema: {sis}</span></td></tr>')

grupos = ['GT · Frente I', 'GT · Frente II', 'GT · Frente III', 'FAETEC', 'Externo']
linhas_ator = ''
for g in grupos:
    linhas_ator += f'<tr class="grp"><td colspan="5">{e(g)}</td></tr>'
    for c, (n, pessoas, grp) in ATORES.items():
        if grp != g or c not in carga: continue
        d = carga[c]
        blocos = []
        for k in ('E', 'X', 'V', 'A'):
            if d[k]: blocos.append(f'<b>{PAPEL[k]} ({len(d[k])}):</b> ' + e(' · '.join(d[k])))
        linhas_ator += (f'<tr><td><b>{e(n)}</b><br><span class="obs">{e(pessoas)}</span></td><td>{"<br>".join(blocos)}</td>'
                        f'<td class="n">{len(d["E"]) + len(d["X"])}</td><td class="n">{d["hoje"]}</td><td class="n">{d["proj"]}</td></tr>')

# matriz compacta 32 × atores
COLS = ['COORD', 'APROV', 'CONF', 'ELAB', 'ATAS', 'PCA', 'TI', 'DOCS', 'DEMAND', 'DGI', 'DIVPL', 'VPA', 'FAETEC', 'ORDEN', 'JUR', 'UCI', 'GEREN', 'FORN', 'SEPLAG', 'PRODERJ', 'CGE']
CURTO = {'COORD': 'Coord.', 'APROV': 'Aprov.', 'CONF': 'Conf.', 'ELAB': 'Elab. (7)', 'ATAS': 'Atas TI', 'PCA': 'PCA', 'TI': 'TI', 'DOCS': 'Docs', 'DEMAND': 'Demand.', 'DGI': 'DGI', 'DIVPL': 'Div. Plan.', 'VPA': 'VPA',
         'FAETEC': 'FAETEC a nomear', 'ORDEN': 'Orden.', 'JUR': 'Jurídico', 'UCI': 'UCI', 'GEREN': 'Gerenc.', 'FORN': 'Fornec.', 'SEPLAG': 'SEPLAG', 'PRODERJ': 'PRODERJ', 'CGE': 'CGE'}
mat = '<tr><th>Atividade</th>' + ''.join(f'<th class="v"><div>{CURTO[c]}</div></th>' for c in COLS) + '</tr>'
for (et, o, nome, p, *_r) in ATIV:
    mat += f'<tr><td class="nm">{ETAPA[et][0]}.{o} {e(nome)}</td>' + ''.join(f'<td class="m {p.get(c, "").lower()}">{p.get(c, "")}</td>' for c in COLS) + '</tr>'

doc = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Atividades por ator · GT PROPAG</title>
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Source+Serif+4:opsz,wght@8..60,600;8..60,700&display=swap" rel="stylesheet">
<style>
@page {{ size: A4 landscape; margin: 13mm 12mm 15mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: Inter, sans-serif; color: #1b2236; font-size: 8.6pt; line-height: 1.45; margin: 0; }}
h1 {{ font-family: 'Source Serif 4', Georgia, serif; color: #fff; font-size: 24pt; line-height: 1.1; margin: 2mm 0 2mm; }}
h2 {{ font-family: 'Source Serif 4', Georgia, serif; color: #16284d; font-size: 14pt; margin: 6mm 0 2mm; padding-top: 2mm; border-top: 1.2px solid #c8962f; break-after: avoid; }}
.quebra {{ break-before: page; }}
.capa {{ background: #16284d; color: #cfd6e6; margin: -13mm -12mm 5mm; padding: 12mm 12mm 8mm; border-left: 9mm solid #c8962f; }}
.eyebrow {{ color: #e0b45a; font-size: 8pt; letter-spacing: .16em; text-transform: uppercase; font-weight: 700; }}
.capa p {{ margin: 1mm 0; font-size: 9.5pt; }}
.regra {{ background: #fbe2de; border-left: 4px solid #b3372b; padding: 2.5mm 4mm; margin: 3mm 0; font-weight: 500; }}
.box {{ background: #fbfaf6; border: 1px solid #e2dccf; border-radius: 2mm; padding: 2.5mm 4mm; margin: 2.5mm 0; }}
table {{ width: 100%; border-collapse: collapse; margin: 2mm 0; }}
thead {{ display: table-header-group; }} tr {{ break-inside: avoid; }}
th {{ background: #16284d; color: #fbfaf6; text-align: left; padding: 1.5mm 2mm; font-weight: 600; font-size: 8pt; }}
td {{ padding: 1.4mm 2mm; border-bottom: 1px solid #e2dccf; vertical-align: top; }}
tr.grp td {{ background: #e1e8f5; color: #16284d; font-weight: 700; font-size: 9pt; padding-top: 2mm; }}
tr.grp td span {{ font-weight: 400; color: #4a4538; font-size: 8pt; }}
.id {{ font-weight: 700; color: #8a5f0d; white-space: nowrap; }}
.obs {{ color: #5f5a4e; font-size: 7.8pt; }}
.alerta {{ color: #b3372b; font-size: 7.8pt; }}
.tag {{ display: inline-block; border-radius: 1mm; padding: .2mm 1.6mm; margin: .3mm 0; font-size: 7.8pt; font-weight: 600; }}
.tag.e {{ background: #16284d; color: #fff; }} .tag.a {{ background: #f3f0e9; color: #4a4538; border: 1px solid #e2dccf; }}
.tag.v {{ background: #f7e9c9; color: #8a5f0d; }} .tag.x {{ background: #fbe2de; color: #b3372b; }}
.base {{ display: inline-block; font-size: 7.4pt; font-weight: 700; padding: .2mm 1.5mm; border-radius: 1mm; }}
.base.doc {{ background: #e2efe6; color: #2f6b4f; }} .base.fun {{ background: #e1e8f5; color: #16284d; }}
.base.nat {{ background: #f3f0e9; color: #5f5a4e; }} .base.prop {{ background: #f7e9c9; color: #8a5f0d; }}
td.n {{ text-align: right; font-family: 'Source Serif 4', Georgia, serif; font-weight: 700; font-size: 11pt; color: #16284d; }}
table.mat th.v {{ height: 21mm; vertical-align: bottom; padding: 1mm .5mm; }}
table.mat th.v div {{ writing-mode: vertical-rl; transform: rotate(180deg); font-size: 7.4pt; white-space: nowrap; }}
table.mat td {{ padding: .55mm 1mm; font-size: 7.2pt; line-height: 1.25; }} table.mat td.nm {{ white-space: nowrap; }}
td.m {{ text-align: center; font-weight: 700; border-left: 1px solid #efe9dc; }}
td.m.e {{ background: #16284d; color: #fff; }} td.m.a {{ background: #f3f0e9; color: #4a4538; }}
td.m.v {{ background: #f7e9c9; color: #8a5f0d; }} td.m.x {{ background: #fbe2de; color: #b3372b; }}
ul {{ margin: 1mm 0 1mm 5mm; padding: 0; }} li {{ margin: .8mm 0; }}
</style></head><body>
<div class="capa">
  <div class="eyebrow">GT PROPAG · Compras Públicas · FAETEC</div>
  <h1>Atividades do fluxo de contratação por ator</h1>
  <p>As 32 atividades das 5 etapas, com quem executa, quem apoia, quem valida e quais órgãos externos respondem. Inclui a carga de cada ator no sistema hoje e na carteira atualizada.</p>
  <p><b style="color:#fff">Posição em 25/09/2026</b> · Complemento da apresentação “Status das contratações da FAETEC” (v2)</p>
</div>
<div class="regra">Nada foi alterado no sistema. Esta matriz é uma proposta de distribuição para validação da coordenação. Designar os responsáveis no sistema depende de autorização.</div>
<div class="box"><b>Como ler.</b>
 Papel na atividade: {cel('ELAB', 'E').replace('Elaboração de artefatos', 'Executa (quem faz)')} {cel('ELAB', 'A').replace('Elaboração de artefatos', 'Apoia / fornece insumo')} {cel('ELAB', 'V').replace('Elaboração de artefatos', 'Valida / aprova')} {cel('ELAB', 'X').replace('Elaboração de artefatos', 'Órgão externo que responde')}.<br>
 Base da atribuição: <span class="base doc">Documento do GT</span> a equipe e o executor da etapa estão no documento “GT FAETEC — Macroprocesso e Equipe”. <span class="base fun">Função no documento</span> a atividade corresponde à função descrita para o integrante. <span class="base nat">Natureza do ato</span> quem faz decorre da norma ou do tipo de ato (ex.: parecer jurídico, empenho). <span class="base prop">Proposta</span> o documento não define quem faz, e a matriz sugere um responsável.</div>

<h2>1. Os atores</h2>
<table><thead><tr><th style="width:22%">Ator</th><th style="width:48%">Quem é</th><th>Grupo</th></tr></thead><tbody>
{''.join(f'<tr><td><b>{e(n)}</b></td><td>{e(p)}</td><td>{e(g)}</td></tr>' for c, (n, p, g) in ATORES.items())}
</tbody></table>

<h2 class="quebra">2. As 32 atividades e quem faz</h2>
<table><thead><tr><th style="width:4%">#</th><th style="width:27%">Atividade · entrega · atenção</th><th style="width:17%">Executa (quem faz)</th><th style="width:19%">Apoia</th><th style="width:17%">Valida · externo</th><th style="width:16%">Base · no sistema</th></tr></thead><tbody>
{linhas_ativ}
</tbody></table>

<h2 class="quebra">3. Por ator: o que cada um faz</h2>
<p>Volume = instâncias das atividades que o ator executa (ou, no caso de órgão externo, a que responde), somadas nas demandas. <b>Hoje</b> = sistema com 17 demandas (276 atividades). <b>Carteira atualizada</b> = 20 da planilha de 24/09 + Tela interativa (354 atividades, projeção).</p>
<table><thead><tr><th style="width:22%">Ator</th><th>Atividades</th><th style="width:7%">Executa</th><th style="width:8%">Volume hoje</th><th style="width:10%">Volume na carteira atualizada</th></tr></thead><tbody>
{linhas_ator}
</tbody></table>

<h2 class="quebra">4. Matriz resumida</h2>
<table class="mat"><thead>{mat}</thead></table>
<p class="obs">E = executa · A = apoia · V = valida/aprova · X = órgão externo que responde. “Elab. (7)” = os 7 integrantes da Frente II.</p>

<h2>5. O que a matriz mostra</h2>
<ul>
 <li><b>As Etapas IV e V não têm dono no GT.</b> São 16 das 32 atividades, 100% FAETEC e sem ponto focal nomeado na Relação de Participantes. Só o ateste técnico tem um integrante do GT proposto (Gibson). Na carteira atualizada, são 144 atividades (16 × 9 demandas com via definida). Nomear esse ponto focal é a lacuna mais relevante da matriz.</li>
 <li><b>A Frente II concentra a Etapa I.</b> Os 7 integrantes executam DOD, ETP, TR e mapa de riscos: 68 atividades hoje e 84 na carteira atualizada, mais as 4 atividades de adesão da Etapa II.</li>
 <li><b>Três pessoas acumulam papéis críticos.</b> Gibson (TI) apoia o PCA, o ETP, o TR e a compatibilidade, e executa a análise do PRODERJ e o ateste técnico. Elias (conformidade) valida 7 atividades, inclusive o gate de ata válida. Cristiane (aprovação) valida 6.</li>
 <li><b>Divergência com o sistema.</b> O documento aponta Jhonatan como ponto focal do PCA, mas no sistema a atividade “Inclusão no PCA / PEDTIC” está com Luene (1 demanda). Vale confirmar antes de designar.</li>
 <li><b>O sistema ainda não reflete a matriz.</b> Só 3 das 276 atividades têm responsável (Luene 1, Gibson 2), e nenhuma demanda tem responsável geral.</li>
 <li><b>Dependência externa nos prazos legais.</b> SEPLAG, PRODERJ e CGE respondem a 3 atividades com prazo legal: 27 na carteira atualizada.</li>
</ul>
<p class="obs">Fontes: documento “GT FAETEC — Macroprocesso e Equipe” (equipe por etapa e funções dos 15 integrantes); sistema em produção, consulta somente leitura de 25/09/2026 (responsáveis designados e instâncias por atividade).</p>
</body></html>"""
(D / 'matriz_atividades_atores.html').write_text(doc)

# ------------------------------------------------------------------ Excel
wb = openpyxl.Workbook()
nav = PatternFill('solid', fgColor='16284D'); cab = Font(bold=True, color='FBFAF6'); fino = Side(style='thin', color='E2DCCF')
def cabecalho(ws, cols, larg):
    ws.append(cols)
    for i, c in enumerate(ws[1], 1):
        c.fill = nav; c.font = cab; c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = larg[i - 1]
    ws.freeze_panes = 'A2'
ws = wb.active; ws.title = 'Atividades x atores'
cabecalho(ws, ['Etapa', 'Nº', 'Atividade', 'Executa (quem faz)', 'Apoia', 'Valida / aprova', 'Externo que responde', 'Entrega', 'Atenção / prazo legal', 'Base da atribuição', 'Observação', 'No sistema (responsável)', 'Instâncias hoje', 'Instâncias carteira atualizada'],
          [8, 5, 32, 34, 34, 30, 24, 30, 36, 20, 44, 20, 10, 12])
for (et, o, nome, p, entrega, alerta, base, obs) in ATIV:
    nm = lambda k: '; '.join(f'{ATORES[c][0]} ({ATORES[c][1]})' if ATORES[c][2].startswith('GT') else ATORES[c][0] for c in quem(p, k))
    ws.append([ETAPA[et][0], f'{ETAPA[et][0]}.{o}', nome, nm('E'), nm('A'), nm('V'), nm('X'), entrega, alerta, BASE[base], obs, NO_SISTEMA.get((et, o), f'0 de {INST_HOJE[et]}'), INST_HOJE[et], INST_PROJ[et]])
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True); c.border = Border(bottom=fino)
ws2 = wb.create_sheet('Por ator')
cabecalho(ws2, ['Ator', 'Quem é', 'Grupo', 'Executa', 'Valida / aprova', 'Apoia', 'Responde (externo)', 'Qtd. executa', 'Volume hoje', 'Volume carteira atualizada'], [28, 50, 16, 60, 40, 50, 34, 10, 10, 12])
for c, (n, pessoas, g) in ATORES.items():
    d = carga.get(c)
    if not d: continue
    ws2.append([n, pessoas, g, '\n'.join(d['E']), '\n'.join(d['V']), '\n'.join(d['A']), '\n'.join(d['X']), len(d['E']) + len(d['X']), d['hoje'], d['proj']])
for row in ws2.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(vertical='top', wrap_text=True); c.border = Border(bottom=fino)
ws3 = wb.create_sheet('Matriz')
cabecalho(ws3, ['Atividade'] + [CURTO[c] for c in COLS], [38] + [9] * len(COLS))
cores = {'E': ('16284D', 'FFFFFF'), 'A': ('F3F0E9', '4A4538'), 'V': ('F7E9C9', '8A5F0D'), 'X': ('FBE2DE', 'B3372B')}
for (et, o, nome, p, *_r) in ATIV:
    ws3.append([f'{ETAPA[et][0]}.{o} {nome}'] + [p.get(c, '') for c in COLS])
    for c in ws3[ws3.max_row][1:]:
        if c.value:
            f, t = cores[c.value]; c.fill = PatternFill('solid', fgColor=f); c.font = Font(bold=True, color=t)
        c.alignment = Alignment(horizontal='center')
ws3.append([]); ws3.append(['E = executa · A = apoia · V = valida/aprova · X = órgão externo que responde'])
wb.save(D / 'Matriz_Atividades_Atores_GT_PROPAG.xlsx')

for c in ('ELAB', 'TI', 'CONF', 'APROV', 'FAETEC', 'DOCS', 'PCA'):
    d = carga[c]; print(c, 'E', len(d['E']), 'A', len(d['A']), 'V', len(d['V']), 'hoje', d['hoje'], 'proj', d['proj'])
print('ok')
